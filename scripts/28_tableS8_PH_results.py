# 28_tableS8_PH_results.py —— Table S8: 比例风险假设检验、时变系数模型与分段敏感性分析
# 输入: raw_data/02_TCGA_CESC_clinical_LONP1_survival.csv, raw_data/44_GSE44001_LONP1_DFS.csv
# 输出: raw_data/52_tableS8_PH_results.csv
#       figures/supplementary files/Additional_file_1_Supplementary_Tables.xlsx 中新增 "Table S8"
# 用法: python3 28_tableS8_PH_results.py <分析根目录>
import sys, os, shutil, warnings
warnings.filterwarnings("ignore")
import numpy as np, pandas as pd
from lifelines import CoxPHFitter, CoxTimeVaryingFitter
from lifelines.statistics import proportional_hazard_test
from lifelines.utils import to_episodic_format
from scipy.stats import chi2 as chi2d

base = sys.argv[1]; rd = os.path.join(base, "raw_data")
d2 = pd.read_csv(os.path.join(rd, "02_TCGA_CESC_clinical_LONP1_survival.csv"))
g44 = pd.read_csv(os.path.join(rd, "44_GSE44001_LONP1_DFS.csv"))
R = []
def add(**kw):
    row = {c: kw.get(c, "") for c in
           ["analysis","cohort","endpoint","model","term","time_months","estimate",
            "lower_95CI","upper_95CI","p_value","n","events","n_at_risk","cumulative_events"]}
    R.append(row)

def prep(df, tk, ek, xcol):
    x = df.dropna(subset=[tk, ek, xcol]); x = x[x[tk] > 0].copy()
    out = x[[tk, ek, xcol]].rename(columns={xcol: "lon"}).reset_index(drop=True)
    out["lon"] = out["lon"] - out["lon"].mean()
    return out

SETS = [("TCGA-CESC","Overall survival", d2,"OS_months","OS_event","LONP1_log2RSEM"),
        ("TCGA-CESC","Disease-specific survival", d2,"DSS_months","DSS_event","LONP1_log2RSEM"),
        ("TCGA-CESC","Progression-free survival", d2,"PFS_months","PFS_event","LONP1_log2RSEM"),
        ("GSE44001","Disease-free survival", g44,"DFS_months","DFS_event","LONP1_log2")]

for coh, ep, src, tk, ek, xc in SETS:
    df = prep(src, tk, ek, xc)
    n, nev = len(df), int(df[ek].sum())
    cf = CoxPHFitter().fit(df, tk, ek)
    s0 = cf.summary.loc["lon"]
    add(analysis="Proportional-hazards (PH) model", cohort=coh, endpoint=ep,
        model="Univariate Cox, LONP1 continuous", term="LONP1 (per log2 unit)",
        estimate=round(np.exp(s0["coef"]),3), lower_95CI=round(np.exp(s0["coef lower 95%"]),3),
        upper_95CI=round(np.exp(s0["coef upper 95%"]),3), p_value=round(s0["p"],4), n=n, events=nev)
    for tr in ("rank","km"):
        r = proportional_hazard_test(cf, df, time_transform=tr).summary
        add(analysis="Schoenfeld residual test of PH", cohort=coh, endpoint=ep,
            model=f"Univariate Cox, {tr} time transform", term="LONP1 (per log2 unit)",
            estimate=round(r.test_statistic.iloc[0],3), p_value=round(r.p.iloc[0],4),
            n=n, events=nev)
    # time-varying
    dd = df.copy(); dd["id"] = np.arange(len(dd))
    long = to_episodic_format(dd, duration_col=tk, event_col=ek, id_col="id")
    long = long[long["stop"] > 0].copy()
    C = float(np.log(df.loc[df[ek]==1, tk]).mean())
    long["gt"] = np.log(long["stop"]) - C
    long["lon_x_gt"] = long["lon"]*long["gt"]
    ctv = CoxTimeVaryingFitter().fit(long[["id","start","stop",ek,"lon","lon_x_gt"]],
          id_col="id", event_col=ek, start_col="start", stop_col="stop", show_progress=False)
    s = ctv.summary; V = ctv.variance_matrix_
    b0, b1 = s.loc["lon","coef"], s.loc["lon_x_gt","coef"]
    add(analysis="Time-varying coefficient model", cohort=coh, endpoint=ep,
        model=f"log HR(t) = b0 + b1 x (log t - c), c = {C:.3f} (t = {np.exp(C):.1f} months)",
        term="b1 (LONP1 x log time): test of time-dependence",
        estimate=round(b1,4), p_value=round(s.loc["lon_x_gt","p"],4), n=n, events=nev)
    for t in [6,12,24,36,60]:
        g = np.log(t)-C; est = b0+b1*g
        var = V.iloc[0,0] + (g**2)*V.iloc[1,1] + 2*g*V.iloc[0,1]
        se = np.sqrt(max(var,0))
        add(analysis="Time-varying coefficient model", cohort=coh, endpoint=ep,
            model="Fitted hazard ratio at time t", term="LONP1 (per log2 unit)", time_months=t,
            estimate=round(np.exp(est),3), lower_95CI=round(np.exp(est-1.96*se),3),
            upper_95CI=round(np.exp(est+1.96*se),3), n=n, events=nev,
            n_at_risk=int((df[tk]>=t).sum()), cumulative_events=int(df.loc[df[tk]<=t, ek].sum()))
    # piecewise
    cut = float(df.loc[df[ek]==1, tk].median())
    l2 = long.copy(); l2["late"] = (l2["stop"]>cut).astype(int)
    l2["lon_early"] = l2["lon"]*(1-l2["late"]); l2["lon_late"] = l2["lon"]*l2["late"]
    c2 = CoxTimeVaryingFitter().fit(l2[["id","start","stop",ek,"lon_early","lon_late"]],
         id_col="id", event_col=ek, start_col="start", stop_col="stop", show_progress=False)
    s2 = c2.summary
    ne = int(df.loc[df[tk]<=cut, ek].sum())
    for nm, tag, nv in [("lon_early", f"t <= {cut:.1f} months", ne), ("lon_late", f"t > {cut:.1f} months", nev-ne)]:
        r = s2.loc[nm]
        add(analysis="Piecewise model (sensitivity analysis)", cohort=coh, endpoint=ep,
            model=f"Split at median event time ({cut:.1f} months)", term=f"LONP1, {tag}",
            estimate=round(np.exp(r["coef"]),3), lower_95CI=round(np.exp(r["coef lower 95%"]),3),
            upper_95CI=round(np.exp(r["coef upper 95%"]),3), p_value=round(r["p"],4), n=n, events=nv)
    # linearity
    from patsy import dmatrix
    ll = cf.log_likelihood_
    B = dmatrix("cr(x, df=4)", {"x": df["lon"].values}, return_type="dataframe").iloc[:,1:]
    B = B.loc[:, B.std()>1e-8]; B.columns=[f"s{i}" for i in range(B.shape[1])]
    dsp = pd.concat([df[[tk,ek]], B], axis=1)
    lls = CoxPHFitter(penalizer=1e-6).fit(dsp, tk, ek).log_likelihood_
    lrt = 2*(lls-ll); ddf = B.shape[1]-1
    add(analysis="Linearity of LONP1 on the log-hazard scale", cohort=coh, endpoint=ep,
        model="Likelihood-ratio test, linear vs natural cubic spline (4 df)",
        term="LONP1 (per log2 unit)", estimate=round(lrt,3),
        p_value=round(chi2d.sf(max(lrt,0), ddf),4), n=n, events=nev)

# multivariable TCGA OS
x = d2.dropna(subset=["OS_months","OS_event"]); x = x[x.OS_months>0].copy()
x["Tg"] = x.path_T.astype(str).str.extract(r"^(T[1-4])")[0]
x = x.dropna(subset=["Tg","age"]); x = x[x.path_N.isin(["N0","N1"])].copy()
x["T2"]=(x.Tg=="T2").astype(int); x["T34"]=x.Tg.isin(["T3","T4"]).astype(int); x["N1"]=(x.path_N=="N1").astype(int)
dm = x[["OS_months","OS_event","LONP1_log2RSEM","age","T2","T34","N1"]].reset_index(drop=True)
cfm = CoxPHFitter().fit(dm,"OS_months","OS_event")
for tr in ("rank","km"):
    r = proportional_hazard_test(cfm, dm, time_transform=tr).summary.reset_index()
    for _, row in r.iterrows():
        add(analysis="Schoenfeld residual test of PH", cohort="TCGA-CESC", endpoint="Overall survival",
            model=f"Multivariable Cox, {tr} time transform", term=row["index"],
            estimate=round(row["test_statistic"],3), p_value=round(row["p"],4),
            n=len(dm), events=int(dm.OS_event.sum()))
    add(analysis="Schoenfeld residual test of PH", cohort="TCGA-CESC", endpoint="Overall survival",
        model=f"Multivariable Cox, {tr} time transform", term="GLOBAL",
        estimate=round(r.test_statistic.sum(),3),
        p_value=round(chi2d.sf(r.test_statistic.sum(), len(r)),4),
        n=len(dm), events=int(dm.OS_event.sum()))
# multivariable GSE44001
xm = g44.dropna(subset=["DFS_months","DFS_event","diameter_cm"]); xm = xm[xm.DFS_months>0].copy()
xm["stage_adv"]=(~xm.stage.isin(["IA1","IA2","IB1"])).astype(int)
dg = xm[["DFS_months","DFS_event","LONP1_log2","stage_adv","diameter_cm"]].reset_index(drop=True)
cfg = CoxPHFitter().fit(dg,"DFS_months","DFS_event")
for tr in ("rank","km"):
    r = proportional_hazard_test(cfg, dg, time_transform=tr).summary.reset_index()
    for _, row in r.iterrows():
        add(analysis="Schoenfeld residual test of PH", cohort="GSE44001", endpoint="Disease-free survival",
            model=f"Multivariable Cox, {tr} time transform", term=row["index"],
            estimate=round(row["test_statistic"],3), p_value=round(row["p"],4),
            n=len(dg), events=int(dg.DFS_event.sum()))
    add(analysis="Schoenfeld residual test of PH", cohort="GSE44001", endpoint="Disease-free survival",
        model=f"Multivariable Cox, {tr} time transform", term="GLOBAL",
        estimate=round(r.test_statistic.sum(),3),
        p_value=round(chi2d.sf(r.test_statistic.sum(), len(r)),4),
        n=len(dg), events=int(dg.DFS_event.sum()))

T = pd.DataFrame(R)
T.to_csv(os.path.join(rd,"52_tableS8_PH_results.csv"), index=False)
print(f"rows={len(T)}")

# ---- write into the supplementary workbook ----
import openpyxl
from openpyxl.styles import Font
xp = os.path.join(base,"figures","supplementary files","Additional_file_1_Supplementary_Tables.xlsx")
bk = os.path.join(base,"figures","supplementary files","_backup_before_TableS8.xlsx")
if not os.path.exists(bk): shutil.copy2(xp, bk); print("backup ->", bk)
wb = openpyxl.load_workbook(xp)
if "Table S8" in wb.sheetnames: del wb["Table S8"]
ws = wb.create_sheet("Table S8")
TNR = Font(name="Times New Roman", size=11)
TNRB = Font(name="Times New Roman", size=11, bold=True)
title = ("Table S8. Proportional-hazards diagnostics for the Cox models of LONP1 expression: Schoenfeld "
         "residual tests, time-varying coefficient models, piecewise sensitivity analyses and linearity checks "
         "(Additional file 2: Fig. S3).")
ws.cell(row=1, column=1, value=title).font = TNR
for j, c in enumerate(T.columns, start=1):
    ws.cell(row=3, column=j, value=c).font = TNRB
for i, (_, row) in enumerate(T.iterrows(), start=4):
    for j, c in enumerate(T.columns, start=1):
        v = row[c]
        ws.cell(row=i, column=j, value=(None if v == "" else v)).font = TNR
for col, w in zip("ABCDEFGHIJKLMN", [34,12,28,46,44,13,11,11,11,10,8,8,10,18]):
    ws.column_dimensions[col].width = w
# README entry
rm = wb["README"]
for r in range(1, 40):
    if rm.cell(row=r, column=1).value and str(rm.cell(row=r, column=1).value).startswith("Table S7"):
        rm.insert_rows(r+1)
        c = rm.cell(row=r+1, column=1,
            value="Table S8  Proportional-hazards diagnostics and time-varying coefficient models for LONP1 survival analyses (Fig. S3)")
        c.font = rm.cell(row=r, column=1).font.copy()
        break
wb.save(xp)
print("Table S8 written to", xp)
